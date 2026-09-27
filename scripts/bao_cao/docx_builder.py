"""Cây node (từ markdown_parser) → một python-docx Document đúng định dạng quy-dinh.md.

Xử lý bẫy OOXML #3 (Bước 3 của quy trình gốc): mọi tiêu đề đầu trang dùng
`quy_dinh.thong_tin.tieu_de_ngan` (ngắn, không xuống 2 dòng ở khổ A4/Times 13pt) thay vì tên
đề tài đầy đủ — tránh tab stop đầu trang bị lệch. Trang bìa nằm ở section đầu tiên, không có
đầu/chân trang; nội dung chính bắt đầu ở section thứ hai với `is_linked_to_previous = False`.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .cau_hinh import QuyDinh
from .markdown_parser import Node
from .numbering import NumberingManager

_BOLD_SPLIT_RE = re.compile(r"(\*\*[^*]+\*\*)")


def _cell_align(a: str) -> WD_ALIGN_PARAGRAPH:
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(a, WD_ALIGN_PARAGRAPH.LEFT)


def _set_table_header_repeat(row) -> None:
    """Bật `w:tblHeader` cho hàng tiêu đề — hàng lặp lại khi bảng tràn trang (Bước 2)."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_repeat_table_header_row_only_first(table) -> None:
    _set_table_header_repeat(table.rows[0])


class BaoCaoBuilder:
    def __init__(self, quy_dinh: QuyDinh, thu_muc_anh: Path):
        self.qd = quy_dinh
        self.thu_muc_anh = thu_muc_anh
        self.document = Document()
        self.numbering = NumberingManager(self.document)
        self.headings: list[tuple[int, str]] = []  # dùng dựng mục lục (muc_luc.py)
        self._da_mo_section_noi_dung = False
        self._can_giua = False  # bật bởi \canle, tắt bởi \hetcanle — dùng cho trang bìa
        self._ap_dung_style_toan_cuc()

    def bat_dau_can_giua(self) -> None:
        self._can_giua = True

    def ket_thuc_can_giua(self) -> None:
        self._can_giua = False

    def them_khoang_trang(self) -> None:
        self.document.add_paragraph()

    # ------------------------------------------------------------------ style toàn cục
    def _ap_dung_style_toan_cuc(self) -> None:
        qd = self.qd.hinh_thuc
        style = self.document.styles["Normal"]
        style.font.name = qd.font_name
        style.font.size = Pt(qd.font_size_pt)
        # Times New Roman cần khai cả rFonts eastAsia/cs để LibreOffice không thay font khác
        # cho ký tự có dấu tiếng Việt.
        rpr = style.element.get_or_add_rPr()
        rFonts = rpr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rpr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), qd.font_name)

        pf = style.paragraph_format
        pf.line_spacing = qd.line_spacing
        pf.space_before = Pt(qd.space_before_pt)
        pf.space_after = Pt(qd.space_after_pt)

        section = self.document.sections[0]
        section.top_margin = Cm(qd.le_tren_cm)
        section.bottom_margin = Cm(qd.le_duoi_cm)
        section.left_margin = Cm(qd.le_trai_cm)
        section.right_margin = Cm(qd.le_phai_cm)

        self._ap_dung_style_heading()

    def _ap_dung_style_heading(self) -> None:
        """Style Heading 1-4 mặc định của python-docx dùng font/màu theo theme (thường là
        sans-serif, màu xanh) — không đúng "Unicode Times New Roman" mà PDF yêu cầu cho toàn
        bộ tài liệu. Ép lại font Times New Roman, màu đen, chỉ khác cỡ chữ/đậm/nghiêng theo
        cấp — đây là lựa chọn trình bày hợp lý cho phân cấp, không phải điều PDF quy định
        cụ thể cỡ chữ tiêu đề."""
        from docx.shared import Pt as _Pt
        from docx.shared import RGBColor as _RGBColor

        qd = self.qd.hinh_thuc
        cau_hinh_cap = {
            1: (16, True, False),
            2: (14, True, False),
            3: (13, True, False),
            4: (13, False, True),
        }
        for cap, (size_pt, dam, nghieng) in cau_hinh_cap.items():
            style = self.document.styles[f"Heading {cap}"]
            style.font.name = qd.font_name
            style.font.size = _Pt(size_pt)
            style.font.bold = dam
            style.font.italic = nghieng
            style.font.color.rgb = _RGBColor(0, 0, 0)
            rpr = style.element.get_or_add_rPr()
            rFonts = rpr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rpr.append(rFonts)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rFonts.set(qn(attr), qd.font_name)
            style.paragraph_format.space_before = _Pt(qd.space_before_pt)
            style.paragraph_format.space_after = _Pt(qd.space_after_pt)

    # ------------------------------------------------------------------ section bìa / nội dung
    def bat_dau_section_noi_dung(self) -> None:
        """Mở section thứ hai (nội dung chính): is_linked_to_previous=False cho header/footer
        — nếu không đặt, Word/LibreOffice nối chân trang ngược lên section bìa (Bước 2)."""
        qd = self.qd.hinh_thuc
        new_section = self.document.add_section(WD_SECTION.NEW_PAGE)
        new_section.top_margin = Cm(qd.le_tren_cm)
        new_section.bottom_margin = Cm(qd.le_duoi_cm)
        new_section.left_margin = Cm(qd.le_trai_cm)
        new_section.right_margin = Cm(qd.le_phai_cm)

        new_section.footer.is_linked_to_previous = False
        new_section.header.is_linked_to_previous = False

        self._dung_dau_trang(new_section)
        self._dung_chan_trang(new_section)
        self._dat_so_trang_bat_dau_lai_tu_1(new_section)
        self._da_mo_section_noi_dung = True

    def _dung_dau_trang(self, section) -> None:
        """Đầu trang chỉ có MỘT phần tử — tiêu đề ngắn `tieu_de_ngan` từ quy-dinh.md, không
        đặt cùng tab stop với phần tử khác. Đây là cách tránh bẫy OOXML #3 (Bước 3): tên đề
        tài dài xuống 2 dòng làm tab stop lệch — vì chỉ có 1 phần tử canh giữa, không có tab
        stop nào để lệch, dù tiêu đề có dài cỡ nào. PDF không bắt buộc phải có đầu trang; đây
        là lựa chọn trình bày, không phải điều khoản của quy định."""
        header_para = section.header.paragraphs[0]
        header_para.text = ""
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_para.add_run(self.qd.thong_tin.tieu_de_ngan)
        run.font.size = Pt(11)
        run.italic = True

    def _dung_chan_trang(self, section) -> None:
        """Chân trang 3 cột qua tab stop: GVHD (trái) — SVTH (giữa) — số trang (phải, đúng
        "góc phải dưới" PDF yêu cầu). GVHD/SVTH là lựa chọn trình bày người dùng chọn thêm,
        số trang ở cực phải là điều khoản bắt buộc của PDF — ưu tiên đúng vị trí này."""
        footer_para = section.footer.paragraphs[0]
        footer_para.text = ""
        pf = footer_para.paragraph_format
        usable_width = (
            section.page_width - section.left_margin - section.right_margin
        )
        pf.tab_stops.add_tab_stop(int(usable_width * 0.5), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)

        gvhd_run = footer_para.add_run(f"GVHD: {self.qd.thong_tin.gvhd}")
        footer_para.add_run("\t")
        svth_run = footer_para.add_run(f"SVTH: {self.qd.thong_tin.svth}")
        footer_para.add_run("\t")
        self._them_truong_so_trang(footer_para)
        for run in (gvhd_run, svth_run):
            run.font.size = Pt(10)

    @staticmethod
    def _them_truong_so_trang(paragraph) -> None:
        run = paragraph.add_run()
        run.font.size = Pt(10)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)

    @staticmethod
    def _dat_so_trang_bat_dau_lai_tu_1(section) -> None:
        """`quy-dinh.md` §3: bắt đầu đánh số từ MỞ ĐẦU = trang 1 — đặt lại `w:pgNumType`
        của section nội dung, không tin giá trị mặc định của Word."""
        sect_pr = section._sectPr
        pg_num_type = sect_pr.find(qn("w:pgNumType"))
        if pg_num_type is None:
            pg_num_type = OxmlElement("w:pgNumType")
            sect_pr.append(pg_num_type)
        pg_num_type.set(qn("w:start"), "1")

    # ------------------------------------------------------------------ nội dung
    def them_doan_van(self, text: str, style: str | None = None) -> None:
        p = self.document.add_paragraph(style=style)
        if self._can_giua:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._them_run_co_dinh_dang(p, text)

    @staticmethod
    def _them_run_co_dinh_dang(paragraph, text: str) -> None:
        """Tách `**đậm**` thành run in đậm riêng — mọi thứ khác (kể cả dấu — em-dash) giữ
        nguyên văn bản, không tự thay thế ký tự."""
        for phan in _BOLD_SPLIT_RE.split(text):
            if not phan:
                continue
            if phan.startswith("**") and phan.endswith("**"):
                run = paragraph.add_run(phan[2:-2])
                run.bold = True
            else:
                paragraph.add_run(phan)

    def them_heading(self, level: int, text: str) -> None:
        p = self.document.add_heading(level=min(level, 4))
        if self._can_giua:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._them_run_co_dinh_dang(p, text)
        self.headings.append((level, text))

    def them_danh_sach_bullet(self, items: list[str]) -> None:
        for item in items:
            p = self.document.add_paragraph(style="List Bullet")
            self._them_run_co_dinh_dang(p, item)

    def them_danh_sach_so_bat_dau_lai(self, items: list[str]) -> None:
        """Mỗi lần gọi tạo MỘT danh sách số độc lập, luôn bắt đầu lại từ 1 (bẫy OOXML #1)."""
        num_id = self.numbering.danh_sach_moi_bat_dau_lai_tu_1()
        for item in items:
            p = self.document.add_paragraph()
            self._them_run_co_dinh_dang(p, item)
            self.numbering.ap_dung_vao_doan(p, num_id)

    def them_bang(self, header: list[str], align: list[str], rows: list[list[str]]) -> None:
        table = self.document.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(header):
            cell = table.rows[0].cells[i]
            cell.paragraphs[0].text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            cell.paragraphs[0].alignment = _cell_align(align[i] if i < len(align) else "left")
        _set_repeat_table_header_row_only_first(table)

        for row_data in rows:
            row = table.add_row()
            for i, text in enumerate(row_data):
                cell = row.cells[i]
                cell.paragraphs[0].text = ""
                self._them_run_co_dinh_dang(cell.paragraphs[0], text)
                cell.paragraphs[0].alignment = _cell_align(align[i] if i < len(align) else "left")
        self.document.add_paragraph()  # đệm giãn cách sau bảng

    def them_hinh(self, duong_dan_anh: Path, caption: str | None, source_line: str | None) -> None:
        from docx.shared import Cm as _Cm

        self.document.add_picture(str(duong_dan_anh), width=_Cm(14))
        last_p = self.document.paragraphs[-1]
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            cap_p = self.document.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap_p.add_run(caption)
            run.bold = True
            run.font.size = Pt(12)
        if source_line:
            src_p = self.document.add_paragraph()
            src_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = src_p.add_run(source_line)
            run.italic = True
            run.font.size = Pt(11)

    def ngat_trang(self) -> None:
        self.document.add_page_break()

    def khoi_ky(self, dong: list[str]) -> None:
        """Khối ký tên không được vỡ đôi khi in — nối `keep_with_next` giữa các dòng liên
        tiếp (bẫy header/footer khác của Bước 3 không áp dụng ở đây, chỉ là yêu cầu Bước 2)."""
        paragraphs = []
        for text in dong:
            p = self.document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._them_run_co_dinh_dang(p, text)
            paragraphs.append(p)
        for p in paragraphs[:-1]:
            p.paragraph_format.keep_with_next = True
        for p in paragraphs:
            p.paragraph_format.keep_together = True

    def them_khoi_ma(self, code: str) -> None:
        """Khối code — font Courier New đơn cách, giữ nguyên xuống dòng, có khung nhẹ để phân
        biệt với văn bản thường (không dùng Table Grid — chỉ tô nền, không cần đường viền đủ 4
        cạnh của một bảng)."""
        p = self.document.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "F2F2F2")
        p.paragraph_format.element.get_or_add_pPr().append(shading)
        for i, dong in enumerate(code.split("\n")):
            if i > 0:
                p.add_run().add_break()
            run = p.add_run(dong if dong else " ")
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            rpr = run._r.get_or_add_rPr()
            rFonts = rpr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rpr.append(rFonts)
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                rFonts.set(qn(attr), "Courier New")

    def them_cong_thuc_van_ban(self, latex: str) -> None:
        """Phương án B cho công thức (Bước 2 mục "công thức"): trình bày dạng văn bản có
        style, dùng khi không render ảnh matplotlib. Giữ nguyên cú pháp LaTeX làm chú thích."""
        p = self.document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(latex.strip())
        run.italic = True

    def luu(self, duong_dan_ra: Path) -> None:
        duong_dan_ra.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(duong_dan_ra))
