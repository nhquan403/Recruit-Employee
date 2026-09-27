"""Dựng mục lục dấu chấm dẫn thật (tab-stop dot leader), không phải chuỗi ký tự `.` gõ tay.

Bước 2 của quy trình gốc đã đo thật: 75 dấu chấm ở Times 13pt chỉ đạt 243,7pt trên 453,5pt bề
chữ — một chuỗi dấu chấm gõ tay là MỘT "từ" liền khối, không co giãn được, dòng sẽ dừng ở đâu
hết chấm chứ không tới đúng lề phải. Cách đúng: `tab_stops.add_tab_stop(pos, RIGHT,
WD_TAB_LEADER.DOTS)` — Word/LibreOffice tự vẽ và co giãn dấu chấm cho vừa đúng khoảng trống.
"""
from __future__ import annotations

from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Pt


def dung_muc_luc(document, section, headings: list[tuple[int, str]], so_trang: dict[str, int] | None = None) -> None:
    """`headings`: danh sách (level, text) đã thu thập lúc build nội dung.
    `so_trang`: map text-tiêu-đề → số trang thật đo được ở lần xuất trước (Phase 9, 2-pass).
    Khi chưa có (lần xuất đầu), in dấu `…` thay vì số — KHÔNG gõ số đoán, chờ đo thật."""
    so_trang = so_trang or {}
    usable_width = section.page_width - section.left_margin - section.right_margin

    heading_p = document.add_paragraph()
    run = heading_p.add_run("MỤC LỤC")
    run.bold = True
    heading_p.alignment = 1  # center

    for level, text in headings:
        p = document.add_paragraph()
        indent_cm = max(0, (level - 1)) * 0.5
        p.paragraph_format.left_indent = Cm(indent_cm)
        p.paragraph_format.tab_stops.add_tab_stop(
            usable_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        run = p.add_run(text)
        if level == 1:
            run.bold = True
        p.add_run("\t")
        trang = so_trang.get(text)
        if trang is not None:
            p.add_run(str(trang))
        elif not so_trang:
            # Lần xuất đầu (chưa đo số trang thật lần nào) — in dấu chờ đo, không gõ số đoán.
            p.add_run("…")
        # Đã đo thật (so_trang khác rỗng) nhưng heading này vẫn không có trong đó: đây là mục
        # thật sự KHÔNG đánh số (VD "MỞ ĐẦU" và các mục con — theo đúng biểu mẫu BM5: "Bắt đầu
        # đánh số trang từ chương 1", mọi trang trước đó không có số trang in ở chân trang) —
        # để trống, không phải "chưa đo xong".


def dung_danh_sach_co_so_trang(
    document, section, muc: list[str], so_trang: dict[str, int] | None = None
) -> None:
    """Danh mục hình/bảng — cùng cơ chế dot-leader như mục lục, nhưng danh sách phẳng (không
    thụt lề theo cấp) vì Hình/Bảng không có cấp con."""
    so_trang = so_trang or {}
    usable_width = section.page_width - section.left_margin - section.right_margin
    for nhan in muc:
        p = document.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(
            usable_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        p.add_run(nhan)
        p.add_run("\t")
        trang = so_trang.get(nhan)
        if trang is not None:
            p.add_run(str(trang))
        elif not so_trang:
            p.add_run("…")
